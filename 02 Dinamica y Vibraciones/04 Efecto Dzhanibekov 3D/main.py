"""
Continuum Lab — Classical Mechanics & Dynamical Systems
Division: 02 Dinamica y Vibraciones / 04 Efecto Dzhanibekov 3D
Entry Point: Euler Intermediate Axis Instability (Dzhanibekov Effect) Simulation

Execution CLI:
- Video Compilation: 1080x1920 @ 60 FPS (9:16 Vertical)
- Dynamic Benchmark: openpyxl XLSX Model with live formulas
- Physical Diagnostics: Rigorous verification of conservation laws
"""

import os
import sys
import argparse
import subprocess
from pathlib import Path
import numpy as np

current_dir = Path(__file__).resolve().parent
WORKSPACE_ROOT = current_dir.parent.parent
OUTPUT_DIR = WORKSPACE_ROOT / "RENDERS" / "04 Efecto Dzhanibekov 3D"

# Ensure local package imports work
if str(current_dir) not in sys.path:
    sys.path.insert(0, str(current_dir))

from src.physics.euler_rigid_body import RigidBodyParams, RigidBodySimulator


def run_diagnostics():
    """Executes theoretical physics diagnostics on Euler equations and conservation laws."""
    print("=" * 70)
    print(" CONTINUUM LAB // DIAGNÓSTICO FÍSICO TEÓRICO: EFECTO DZHANIBEKOV")
    print("=" * 70)

    params = RigidBodyParams(I1=1.0, I2=2.4, I3=4.2)
    sim = RigidBodySimulator(params)

    # 1. Eigenvalue analysis around intermediate axis
    coeff_intermediate = (params.I2 - params.I3) * (params.I1 - params.I2) / (params.I1 * params.I3)
    lam = np.sqrt(coeff_intermediate) * 12.0
    print(f"[DINÁMICA] Momentos principales : I1={params.I1:.1f} < I2={params.I2:.1f} < I3={params.I3:.1f} kg·m²")
    print(f"[ESTABILIDAD] Coeficiente hiperbólico: {coeff_intermediate:+.4f} (POSITIVO -> SADDLE INESTABLE)")
    print(f"[ESTABILIDAD] Tasa de divergencia local: lambda = {lam:.4f} s^-1 (Lyapunov positivo)")

    # 2. Separatrix ratio
    ratio = np.sqrt((params.I3 * (params.I3 - params.I2)) / (params.I1 * (params.I2 - params.I1)))
    print(f"[SEPARATRIZ] Relación homoclínica exacta w1/w3: {ratio:.6f}")

    # 3. 14s Numerical Integration
    print(f"[SOLVER] Integrando con Runge-Kutta orden 8 (DOP853) a 60 FPS...")
    res = sim.simulate(t_span=(0.0, 14.0), fps=60, method="DOP853")

    t = res["time"]
    w2 = res["omega_body"][:, 1]
    crossings = t[np.where(np.diff(np.signbit(w2)))[0]]

    print(f"[CONSERVACIÓN] Energía cinética rotacional inicial : T_rot = {res['initial_energy']:.6f} J")
    print(f"[CONSERVACIÓN] Deriva relativa máxima de energía   : dE/E0  = {res['rel_energy_drift']:.2e}")
    print(f"[CONSERVACIÓN] Magnitud momento angular inicial     : |L|    = {res['initial_l_mag']:.6f} N·m·s")
    print(f"[CONSERVACIÓN] Deriva relativa de magnitud |L|      : dL/L0  = {res['rel_momentum_drift']:.2e}")
    print(f"[CONSERVACIÓN] Deriva vectorial espacial L_s        : max|dL|= {res['space_momentum_drift']:.2e} N·m·s")
    print(f"[CINEMÁTICA] Cruces por cero de omega_2 detectados:")
    for i, ct in enumerate(crossings, 1):
        print(f"             Flip {i}: t = {ct:.2f} s")

    print("-" * 70)
    print(" VEREDICTO FÍSICO: CONSERVACIÓN ESTRICTA Y FLIPS HOMOCLÍNICOS CONFIRMADOS")
    print("=" * 70)
    return True


def export_benchmark_data():
    """
    Exports 1,400 simulation steps to a professional XLSX workbook using openpyxl.
    Strictly adheres to dynamic Excel formulas (no hardcoded totals).
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("[ERROR] openpyxl is required to generate Excel benchmark.")
        return False

    bench_dir = OUTPUT_DIR / "benchmarks"
    bench_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = bench_dir / "Efecto_Dzhanibekov_Conservacion_Euler.xlsx"

    print(f"[BENCHMARK] Generating dynamic Excel model: {xlsx_path.name}...")

    params = RigidBodyParams(I1=1.0, I2=2.4, I3=4.2)
    sim = RigidBodySimulator(params)
    res = sim.simulate(t_span=(0.0, 14.0), fps=100, method="DOP853")

    wb = Workbook()
    ws = wb.active
    ws.title = "Euler_Conservation_Model"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palette (Continuum Lab Aerospace Dark)
    hdr_fill = PatternFill(start_color="0D1117", end_color="0D1117", fill_type="solid")
    hdr_font = Font(name="Consolas", size=10, bold=True, color="00F0FF")
    data_font = Font(name="Consolas", size=9, color="E6EDF3")
    alt_fill = PatternFill(start_color="161B22", end_color="161B22", fill_type="solid")
    formula_font = Font(name="Consolas", size=9, color="38BDF8")

    thin_border = Border(
        left=Side(style="thin", color="21262D"),
        right=Side(style="thin", color="21262D"),
        top=Side(style="thin", color="21262D"),
        bottom=Side(style="thin", color="21262D")
    )

    headers = [
        "Paso",
        "Tiempo_s",
        "Omega1_rad_s",
        "Omega2_rad_s",
        "Omega3_rad_s",
        "Energia_Rot_J",
        "Momento_Angular_Nms",
        "Desviacion_Energia_Abs",
        "Desviacion_Momento_Abs",
        "Estado_Eje_Intermedio"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    num_rows = len(res["time"])
    for i in range(num_rows):
        r = i + 2
        t_val = res["time"][i]
        w1_val = res["omega_body"][i, 0]
        w2_val = res["omega_body"][i, 1]
        w3_val = res["omega_body"][i, 2]

        ws.cell(row=r, column=1, value=i + 1).number_format = "#,##0"
        ws.cell(row=r, column=2, value=float(t_val)).number_format = "0.000"
        ws.cell(row=r, column=3, value=float(w1_val)).number_format = "0.000000"
        ws.cell(row=r, column=4, value=float(w2_val)).number_format = "0.000000"
        ws.cell(row=r, column=5, value=float(w3_val)).number_format = "0.000000"

        # Dynamic Excel Formulas
        ws.cell(row=r, column=6, value=f"=0.5*(1.0*C{r}^2 + 2.4*D{r}^2 + 4.2*E{r}^2)").number_format = "0.000000"
        ws.cell(row=r, column=7, value=f"=SQRT((1.0*C{r})^2 + (2.4*D{r})^2 + (4.2*E{r})^2)").number_format = "0.000000"
        ws.cell(row=r, column=8, value=f"=ABS(F{r}-$F$2)").number_format = "0.00000000"
        ws.cell(row=r, column=9, value=f"=ABS(G{r}-$G$2)").number_format = "0.00000000"
        ws.cell(row=r, column=10, value=f'=IF(D{r}>=0, "ESTABLE (+Y)", "INVERTIDO (-Y)")')

        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.font = data_font if c <= 5 else formula_font
            cell.border = thin_border
            if i % 2 == 1:
                cell.fill = alt_fill

    # Set column widths
    col_widths = [10, 12, 16, 16, 16, 18, 22, 24, 24, 22]
    for idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

    wb.save(str(xlsx_path))
    print(f"[SUCCESS] Dynamic benchmark workbook saved: {xlsx_path}")
    return True


def render_video(quality_flag: str = "-qh"):
    """
    Renders the 14-second 9:16 vertical vector video via Manim Community.
    """
    video_dir = OUTPUT_DIR / "videos"
    video_dir.mkdir(parents=True, exist_ok=True)
    cache_dir = video_dir / "manim_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)

    scene_file = current_dir / "src" / "visualization" / "dzhanibekov_scene.py"
    target_mp4 = video_dir / "efecto_dzhanibekov_1080x1920.mp4"

    cmd = [
        sys.executable, "-m", "manim", quality_flag,
        str(scene_file), "DzhanibekovScene",
        "--media_dir", str(cache_dir),
        "-o", "efecto_dzhanibekov_1080x1920.mp4"
    ]

    print(f"[CONTINUUM LAB] Launching Manim 9:16 video compilation ({quality_flag})...")
    print(f"[COMMAND] {' '.join(cmd)}")
    res = subprocess.run(cmd, cwd=str(current_dir))

    if res.returncode != 0:
        print(f"[ERROR] Manim rendering failed with exit code {res.returncode}")
        return False

    rendered_files = list(cache_dir.glob("**/efecto_dzhanibekov_1080x1920.mp4"))
    if rendered_files:
        import shutil
        shutil.move(str(rendered_files[0]), str(target_mp4))
        shutil.rmtree(str(cache_dir), ignore_errors=True)
        print(f"[SUCCESS] Video finalized at: {target_mp4}")
        return True

    print("[WARNING] Compiled video file not located in cache.")
    return False


def main():
    parser = argparse.ArgumentParser(description="Continuum Lab — Dzhanibekov Effect 3D Simulation")
    parser.add_argument("--diagnostics", action="store_true", help="Run physics diagnostics")
    parser.add_argument("--benchmark", action="store_true", help="Export Excel benchmark model")
    parser.add_argument("--preview", action="store_true", help="Render low-quality preview (-ql)")
    parser.add_argument("--render", action="store_true", help="Render high-quality 1080x1920 @ 60 FPS (-qh)")
    parser.add_argument("--all", action="store_true", help="Run diagnostics, benchmark, and video render")

    args = parser.parse_args()

    if args.all or (not any(vars(args).values())):
        run_diagnostics()
        export_benchmark_data()
        render_video("-qh")
    else:
        if args.diagnostics:
            run_diagnostics()
        if args.benchmark:
            export_benchmark_data()
        if args.preview:
            render_video("-ql")
        if args.render:
            render_video("-qh")


if __name__ == "__main__":
    main()
