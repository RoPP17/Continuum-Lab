"""
Continuum Lab — Classical Mechanics & Dynamical Systems
Division: 02 Dinamica y Vibraciones / 02 Mecanismo Coriolis Collarin
Module: export_coriolis_benchmark.py

Automated Generation of Professional Engineering Excel Model (OpenPyXL)
Fully Dynamic Formulas, Zero Hardcoding, Continuum Lab Visual Standard.
Complies strictly with user-defined XLSX engineering skill rules.
"""

import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Add project root to sys.path
project_root = Path(__file__).resolve().parent.parent.parent
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from src.physics.coriolis_kinematics import CoriolisMechanismParams, CoriolisKinematicsSolver


def export_coriolis_benchmark(output_path: Path = None) -> Path:
    """Generates a professional engineering spreadsheet with dynamic kinematic formulas."""
    if output_path is None:
        workspace_root = Path(__file__).resolve().parent.parent.parent.parent.parent
        output_dir = workspace_root / "RENDERS" / "Coriolis_Mechanism"
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / "Mecanismo_Coriolis_Cinematica_Benchmark.xlsx"

    wb = openpyxl.Workbook()

    # Hoja 1: Modelo Cinemático Dinámico
    ws = wb.active
    ws.title = "Coriolis_Kinematics"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de Estilos Continuum Lab (Tema Ingeniería Oscura / Alta Legibilidad)
    font_family = "Segoe UI"
    hdr_fill = PatternFill(start_color="0D1117", end_color="0D1117", fill_type="solid")
    hdr_font = Font(name=font_family, size=11, bold=True, color="00F0FF")
    subhdr_fill = PatternFill(start_color="161B22", end_color="161B22", fill_type="solid")
    subhdr_font = Font(name=font_family, size=10, bold=True, color="58A6FF")

    input_fill = PatternFill(start_color="0D223A", end_color="0D223A", fill_type="solid")
    input_font = Font(name=font_family, size=10, bold=True, color="79C0FF")

    cell_font = Font(name=font_family, size=10, color="E6EDF3")
    accent_font = Font(name=font_family, size=10, bold=True, color="00F0FF")
    alt_fill = PatternFill(start_color="161B22", end_color="161B22", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="30363D"),
        right=Side(style="thin", color="30363D"),
        top=Side(style="thin", color="30363D"),
        bottom=Side(style="thin", color="30363D"),
    )

    # 1. BLOQUE DE PARÁMETROS DE ENTRADA (CELDAS EDITABLES DE DISEÑO)
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "CONTINUUM LAB • MECANISMO DE RETORNO CON ACELERACION DE CORIOLIS"
    title_cell.font = Font(name=font_family, size=12, bold=True, color="00F0FF")
    title_cell.fill = hdr_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    params_meta = [
        ("Longitud Manivela O1-A (L1)", 1.00, "m", "Longitud de la barra impulsora giratoria", "B3"),
        ("Distancia entre Pivotes (d)", 1.55, "m", "Separacion fija entre ejes O1 y O2", "B4"),
        ("Velocidad Angular Manivela (w1)", 2.50, "rad/s", "Velocidad constante impuesta al motor", "B5"),
        ("Paso de Integracion Temporal (dt)", 0.005, "s", "Delta de tiempo entre muestras", "B6"),
        ("Regimen de Operacion", '=IF(B4>B3,"Oscilante (d > L1)","Rotacion Continua Whitworth")', "-", "Tipo de trayectoria del balancin ranurado", "B7"),
    ]

    ws["A2"] = "PARAMETROS DE ENTRADA"
    ws["B2"] = "VALOR"
    ws["C2"] = "UNIDAD"
    ws["D2"] = "DESCRIPCION TECNICA"
    for col in ["A2", "B2", "C2", "D2"]:
        ws[col].font = subhdr_font
        ws[col].fill = subhdr_fill
        ws[col].border = thin_border

    for idx, (p_name, p_val, p_unit, p_desc, coord) in enumerate(params_meta, start=3):
        ws[f"A{idx}"] = p_name
        cell_val = ws[coord]
        cell_val.value = p_val
        cell_val.font = input_font
        cell_val.fill = input_fill
        cell_val.alignment = Alignment(horizontal="right")
        cell_val.border = thin_border

        if isinstance(p_val, float):
            cell_val.number_format = "0.00" if "dt" not in p_name else "0.000"

        ws[f"C{idx}"] = p_unit
        ws[f"D{idx}"] = p_desc
        for c in [f"A{idx}", f"C{idx}", f"D{idx}"]:
            ws[c].font = cell_font
            ws[c].border = thin_border

    # 2. ENCABEZADOS DE LA TABLA DINÁMICA DE CINEMÁTICA
    data_start_row = 10
    headers = [
        ("Paso", "#,##0"),
        ("Tiempo [s]", "0.000"),
        ("Theta1 [rad]", "0.0000"),
        ("xA [m]", "0.0000"),
        ("yA [m]", "0.0000"),
        ("r2 [m]", "0.0000"),
        ("Theta2 [rad]", "0.0000"),
        ("vxA [m/s]", "0.0000"),
        ("vyA [m/s]", "0.0000"),
        ("v_rel (dr2/dt) [m/s]", "0.0000"),
        ("omega2 [rad/s]", "0.0000"),
        ("a_Coriolis (2*w2*v_rel) [m/s²]", "0.0000"),
        ("a_Centripeta [m/s²]", "0.0000"),
    ]

    for c_idx, (h_title, _) in enumerate(headers, start=1):
        cell = ws.cell(row=data_start_row, column=c_idx, value=h_title)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[data_start_row].height = 28

    # 3. GENERACIÓN DE FILAS CON FÓRMULAS EXCEL DINÁMICAS (SIN HARDCODING)
    # L1: $B$3, d: $B$4, w1: $B$5, dt: $B$6
    num_steps = 400
    for i in range(num_steps):
        r = data_start_row + 1 + i
        step_idx = i

        # Paso
        ws.cell(row=r, column=1, value=step_idx)
        # Tiempo t = Paso * dt
        ws.cell(row=r, column=2, value=f"=A{r}*$B$6")
        # Theta1 = w1 * t
        ws.cell(row=r, column=3, value=f"=$B$5*B{r}")
        # xA = L1 * COS(Theta1)
        ws.cell(row=r, column=4, value=f"=$B$3*COS(C{r})")
        # yA = L1 * SIN(Theta1)
        ws.cell(row=r, column=5, value=f"=$B$3*SIN(C{r})")
        # r2 = SQRT(xA^2 + (yA + d)^2)
        ws.cell(row=r, column=6, value=f"=SQRT(D{r}^2 + (E{r}+$B$4)^2)")
        # Theta2 = ATAN2(xA, yA + d)
        ws.cell(row=r, column=7, value=f"=ATAN2(D{r}, E{r}+$B$4)")
        # vxA = -L1 * w1 * SIN(Theta1)
        ws.cell(row=r, column=8, value=f"=-$B$3*$B$5*SIN(C{r})")
        # vyA = L1 * w1 * COS(Theta1)
        ws.cell(row=r, column=9, value=f"=$B$3*$B$5*COS(C{r})")
        # v_rel = vxA * COS(Theta2) + vyA * SIN(Theta2)
        ws.cell(row=r, column=10, value=f"=H{r}*COS(G{r}) + I{r}*SIN(G{r})")
        # omega2 = (-vxA * SIN(Theta2) + vyA * COS(Theta2)) / r2
        ws.cell(row=r, column=11, value=f"=(-H{r}*SIN(G{r}) + I{r}*COS(G{r}))/F{r}")
        # a_Coriolis = 2 * omega2 * v_rel
        ws.cell(row=r, column=12, value=f"=2*K{r}*J{r}")
        # a_Centripeta = -omega2^2 * r2
        ws.cell(row=r, column=13, value=f"=-K{r}^2*F{r}")

        # Aplicar formatos numéricos y estilos
        for c_idx, (_, num_fmt) in enumerate(headers, start=1):
            c_cell = ws.cell(row=r, column=c_idx)
            c_cell.font = accent_font if c_idx == 12 else cell_font
            c_cell.number_format = num_fmt
            c_cell.border = thin_border
            if i % 2 == 1:
                c_cell.fill = alt_fill

    # 4. RESUMEN ESTADÍSTICO EN LA PARTE INFERIOR
    summary_row = data_start_row + num_steps + 2
    ws.cell(row=summary_row, column=1, value="MÁXIMO ABSOLUTO").font = subhdr_font
    ws.cell(row=summary_row, column=12, value=f"=MAX(L{data_start_row+1}:L{summary_row-2})").number_format = "0.0000"
    ws.cell(row=summary_row, column=12).font = accent_font

    ws.cell(row=summary_row+1, column=1, value="MÍNIMO ABSOLUTO").font = subhdr_font
    ws.cell(row=summary_row+1, column=12, value=f"=MIN(L{data_start_row+1}:L{summary_row-2})").number_format = "0.0000"
    ws.cell(row=summary_row+1, column=12).font = accent_font

    ws.cell(row=summary_row+2, column=1, value="PROMEDIO CUADRÁTICO (RMS)").font = subhdr_font
    ws.cell(row=summary_row+2, column=12, value=f"=SQRT(SUMSQ(L{data_start_row+1}:L{summary_row-2})/{num_steps})").number_format = "0.0000"
    ws.cell(row=summary_row+2, column=12).font = accent_font

    # Autoajuste de ancho de columnas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            v_str = str(cell.value or "")
            if len(v_str) > max_len and not cell.coordinate in ["A1", "B1", "C1", "D1"]:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(str(output_path))
    print(f"[SUCCESS] Engineering Excel model exported to: {output_path}")
    return output_path


if __name__ == "__main__":
    export_coriolis_benchmark()
