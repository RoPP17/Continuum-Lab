"""
Continuum Lab - Unit & Physics Verification Suite for Chaotic Three-Body Problem
Module: 02 Dinamica y Vibraciones / 06 Problema de los Tres Cuerpos Caotico
File: tests/test_three_body.py
"""

import os
import tempfile
import wave
from pathlib import Path
import pytest
import numpy as np
import openpyxl

from src.physics.three_body_sim import (
    ThreeBodyConfig,
    ThreeBodySimulation,
)
from src.physics.export_excel_benchmark import create_three_body_excel_benchmark
from src.audio.three_body_audio import synthesize_three_body_soundtrack


def test_hamiltonian_conservation_periodic_orbit():
    """Verify energy conservation along the Moore-Chenciner Figure-8 orbit."""
    cfg = ThreeBodyConfig(perturb_time=999.0)  # No perturbation for this test
    sim = ThreeBodySimulation(config=cfg)

    telem_init = sim.step(0.001)
    E0 = telem_init.total_energy
    assert E0 < 0.0  # Gravitationally bound system

    # Advance along figure-8 for 3.0 seconds
    telem_3s = sim.advance_to_time(3.0)
    E1 = telem_3s.total_energy

    # Relative energy conservation error should be < 0.5%
    rel_error = abs((E1 - E0) / E0)
    assert rel_error < 0.005


def test_angular_momentum_and_momentum_conservation():
    """Verify that total angular and linear momentum are zero."""
    sim = ThreeBodySimulation()
    telem = sim.step(0.001)

    assert np.isclose(telem.angular_momentum, 0.0, atol=1e-3)

    # Linear momentum
    px = sim.state[6] + sim.state[8] + sim.state[10]
    py = sim.state[7] + sim.state[9] + sim.state[11]
    assert np.isclose(px, 0.0, atol=1e-4)
    assert np.isclose(py, 0.0, atol=1e-4)


def test_chaotic_divergence_and_slingshots():
    """Verify that perturbation triggers chaotic divergence and velocity spikes."""
    cfg = ThreeBodyConfig(perturb_time=2.0, perturb_dx=1e-4, perturb_dy=-1e-4)
    sim = ThreeBodySimulation(config=cfg)

    # Before perturbation
    telem_early = sim.advance_to_time(1.0)
    assert not telem_early.is_chaotic

    # After perturbation and chaotic close encounter
    telem_late = sim.advance_to_time(6.0)
    assert telem_late.is_chaotic
    assert telem_late.max_speed > 0.0
    assert telem_late.min_separation > 0.0


def test_excel_benchmark_generation():
    """Verify dynamic Excel benchmark generation and dynamic formula structure."""
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = os.path.join(tmpdir, "test_3body_benchmark.xlsx")
        create_three_body_excel_benchmark(xlsx_path)

        assert os.path.exists(xlsx_path)
        wb = openpyxl.load_workbook(xlsx_path)
        assert "Parametros_y_Conservacion" in wb.sheetnames
        assert "Taxonomia_Soluciones_3Cuerpos" in wb.sheetnames

        ws1 = wb["Parametros_y_Conservacion"]
        assert ws1["A5"].value.startswith("=")
        assert "SQRT" in ws1["G9"].value


def test_audio_soundtrack_synthesis():
    """Verify procedural audio WAV synthesis."""
    with tempfile.TemporaryDirectory() as tmpdir:
        wav_path = os.path.join(tmpdir, "test_3body_audio.wav")
        synthesize_three_body_soundtrack(wav_path, duration=0.5, sample_rate=22050)

        assert os.path.exists(wav_path)
        with wave.open(wav_path, "rb") as wf:
            assert wf.getnchannels() == 2
            assert wf.getsampwidth() == 2
            assert wf.getframerate() == 22050
            frames = wf.readframes(wf.getnframes())
            data = np.frombuffer(frames, dtype=np.int16)
            assert len(data) > 0
            assert np.max(np.abs(data)) > 1000
