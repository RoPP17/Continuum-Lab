"""
Continuum Lab - Test Suite for Rayleigh-Taylor Instability Suite (#21)
Module: 01 Mecanica de Fluidos / 06 Inestabilidad Rayleigh Taylor
File: tests/test_rayleigh_taylor.py
"""

import os
import sys
import wave
import pytest
import numpy as np
import openpyxl

# Add module and project root to sys.path
test_dir = os.path.dirname(os.path.abspath(__file__))
module_root = os.path.dirname(test_dir)
project_root = os.path.dirname(os.path.dirname(module_root))

if module_root not in sys.path:
    sys.path.insert(0, module_root)
if project_root not in sys.path:
    sys.path.insert(0, project_root)

from src.physics.rayleigh_taylor_sim import (
    RayleighTaylorConfig,
    RayleighTaylorSimulation,
)
from src.physics.export_excel_benchmark import create_rayleigh_taylor_excel_benchmark
from src.audio.rayleigh_taylor_audio import synthesize_rayleigh_taylor_soundtrack


def test_simulation_initialization():
    """Verifies grid dimensions, bounds, and initial conditions."""
    cfg = RayleighTaylorConfig(Nx=64, Ny=128, duration=1.0)
    sim = RayleighTaylorSimulation(cfg)

    assert sim.X.shape == (128, 64)
    assert sim.Y.shape == (128, 64)
    assert np.all(sim.C >= 0.0) and np.all(sim.C <= 1.0)
    assert np.max(np.abs(sim.omega)) == 0.0
    assert len(sim.telemetry_history) == 1
    assert sim.telemetry_history[0].atwood_number == 0.50


def test_incompressibility():
    """Verifies that the velocity field is strictly divergence-free (∇·u ≈ 0)."""
    cfg = RayleighTaylorConfig(Nx=64, Ny=128, duration=1.0)
    sim = RayleighTaylorSimulation(cfg)

    # Run a few steps to generate motion
    for _ in range(10):
        sim.step()

    # Compute divergence in spectral space: 1j*KX*u_hat + 1j*KY*v_hat
    u_hat = np.fft.fft2(sim.u)
    v_hat = np.fft.fft2(sim.v)
    div_hat = 1j * sim.KX * u_hat + 1j * sim.KY * v_hat
    div = np.real(np.fft.ifft2(div_hat))

    assert np.max(np.abs(div)) < 1e-8


def test_concentration_conservation():
    """Verifies that total concentration / fluid mass is conserved."""
    cfg = RayleighTaylorConfig(Nx=64, Ny=128, duration=1.0)
    sim = RayleighTaylorSimulation(cfg)

    initial_mass = np.sum(sim.C) * sim.dx * sim.dy
    for _ in range(25):
        sim.step()
    final_mass = np.sum(sim.C) * sim.dx * sim.dy

    rel_error = abs(final_mass - initial_mass) / initial_mass
    assert rel_error < 0.02  # Mass conserved within 2% across sponge/diffusion


def test_baroclinic_torque_generation():
    """Verifies that horizontal density gradient creates baroclinic vorticity."""
    cfg = RayleighTaylorConfig(Nx=64, Ny=128, duration=1.0)
    sim = RayleighTaylorSimulation(cfg)

    assert np.max(np.abs(sim.omega)) == 0.0
    sim.step()
    # After one step, baroclinic torque must have generated non-zero vorticity
    assert np.max(np.abs(sim.omega)) > 0.0


def test_mixing_layer_growth():
    """Verifies that the mixing layer grows as heavy fluid descends."""
    cfg = RayleighTaylorConfig(Nx=64, Ny=128, duration=1.0)
    sim = RayleighTaylorSimulation(cfg)

    initial_h = sim.telemetry_history[0].mixing_width
    for _ in range(40):
        sim.step()
    sim._record_telemetry()
    final_h = sim.telemetry_history[-1].mixing_width

    assert final_h >= initial_h


def test_audio_generation(tmp_path):
    """Verifies 44.1 kHz stereo audio generation and peak constraints."""
    out_wav = str(tmp_path / "test_audio.wav")
    synthesize_rayleigh_taylor_soundtrack(out_wav, duration=1.5, sample_rate=44100)

    assert os.path.exists(out_wav)
    with wave.open(out_wav, "rb") as wf:
        assert wf.getnchannels() == 2
        assert wf.getsampwidth() == 2
        assert wf.getframerate() == 44100
        n_frames = wf.getnframes()
        assert n_frames == int(1.5 * 44100)
        raw = wf.readframes(n_frames)
        samples = np.frombuffer(raw, dtype=np.int16)
        assert len(samples) > 0
        assert np.max(np.abs(samples)) > 1000
        assert np.max(np.abs(samples)) <= 32767


def test_excel_benchmark_generation(tmp_path):
    """Verifies Excel generation with dynamic formulas and correct sheet structure."""
    out_xlsx = str(tmp_path / "test_benchmark.xlsx")
    create_rayleigh_taylor_excel_benchmark(out_xlsx)

    assert os.path.exists(out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx, data_only=False)

    expected_sheets = [
        "Parametros_Fisicos",
        "Dispersion_y_Crecimiento",
        "Evolucion_Temporal",
        "Comparativa_Sistemas",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames
        ws = wb[s]
        assert ws.max_row > 5
        assert ws.max_column >= 5

    # Check that formulas exist and start with '='
    ws1 = wb["Parametros_Fisicos"]
    formula_count = 0
    for row in ws1.iter_rows(values_only=True):
        for val in row:
            if val is not None and str(val).startswith("="):
                formula_count += 1
                assert str(val) == str(val).upper()  # Uppercase formulas
    assert formula_count >= 5
